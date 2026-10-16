import io
import os
from typing import List, Dict
import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START, WD_ORIENT

from src.parsers.base import DocumentChunk, ElementType
from .base import BaseBuilder, align_chunk_translations

DEFAULT_PAGE_WIDTH = Inches(8.27)
DEFAULT_PAGE_HEIGHT = Inches(11.69)
DEFAULT_MARGIN = Inches(1)

def _apply_text_section_format(section):
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = DEFAULT_PAGE_WIDTH
    section.page_height = DEFAULT_PAGE_HEIGHT
    section.top_margin = DEFAULT_MARGIN
    section.bottom_margin = DEFAULT_MARGIN
    section.left_margin = DEFAULT_MARGIN
    section.right_margin = DEFAULT_MARGIN

class DocxBuilder(BaseBuilder):
    def build(
        self,
        chunks: List[DocumentChunk],
        translations: Dict[int, Dict[int, str]],
        output_file_path: str,
        bilingual: bool = False
    ):
        doc = docx.Document()

        # Atur format section awal untuk teks standar
        if doc.sections:
            _apply_text_section_format(doc.sections[0])

        # Atur Font & Format Default (Normal Style)
        style = doc.styles['Normal']
        font = style.font
        font.name = 'Times New Roman'
        font.size = Pt(12)
        style.paragraph_format.line_spacing = 1.15
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.first_line_indent = Inches(0.35)
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

        in_image_section = False
        has_any_content = False

        for chunk in chunks:
            chunk_trans = align_chunk_translations(chunk, translations.get(chunk.chunk_id, {}))

            for i, el in enumerate(chunk.elements):
                if el.element_type == ElementType.PAGE_BREAK:
                    if not in_image_section:
                        doc.add_page_break()
                    continue

                if el.is_image() and el.image_bytes:
                    try:
                        from PIL import Image
                        img_stream = io.BytesIO(el.image_bytes)
                        pil_img = Image.open(img_stream).convert("RGB")
                        img_w, img_h = pil_img.size

                        clean_stream = io.BytesIO()
                        pil_img.save(clean_stream, format="JPEG")
                        clean_stream.seek(0)

                        if img_w >= 200 and img_h >= 200:
                            # Gambar ilustrasi/cover: buat section presisi sesuai ukuran & rasio gambar
                            if not has_any_content:
                                sec = doc.sections[0]
                            else:
                                sec = doc.add_section(WD_SECTION_START.NEW_PAGE)

                            sec.top_margin = Inches(0)
                            sec.bottom_margin = Inches(0)
                            sec.left_margin = Inches(0)
                            sec.right_margin = Inches(0)

                            if img_w > img_h:
                                # Gambar Landscape (spread horizontal)
                                target_w = DEFAULT_PAGE_HEIGHT
                                target_h = Inches(DEFAULT_PAGE_HEIGHT.inches * (img_h / img_w))
                                sec.orientation = WD_ORIENT.LANDSCAPE
                                sec.page_width = target_w
                                sec.page_height = target_h
                            else:
                                # Gambar Portrait (vertikal)
                                target_h = DEFAULT_PAGE_HEIGHT
                                target_w = Inches(DEFAULT_PAGE_HEIGHT.inches * (img_w / img_h))
                                sec.orientation = WD_ORIENT.PORTRAIT
                                sec.page_width = target_w
                                sec.page_height = target_h

                            p = doc.paragraphs[0] if (not has_any_content and doc.paragraphs) else doc.add_paragraph()
                            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            p.paragraph_format.first_line_indent = Inches(0)
                            p.paragraph_format.space_before = Pt(0)
                            p.paragraph_format.space_after = Pt(0)
                            p.paragraph_format.line_spacing = Pt(1)
                            run = p.add_run()
                            run.add_picture(clean_stream, width=target_w, height=target_h - Pt(4))
                            in_image_section = True
                            has_any_content = True
                        else:
                            # Gambar kecil (inline icon / ornament)
                            if in_image_section:
                                sec = doc.add_section(WD_SECTION_START.NEW_PAGE)
                                _apply_text_section_format(sec)
                                in_image_section = False

                            p = doc.add_paragraph()
                            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            p.paragraph_format.first_line_indent = Inches(0)
                            run = p.add_run()
                            run.add_picture(clean_stream, width=Inches(min(5.2, img_w / 96.0)))
                            has_any_content = True
                    except Exception:
                        if in_image_section:
                            sec = doc.add_section(WD_SECTION_START.NEW_PAGE)
                            _apply_text_section_format(sec)
                            in_image_section = False
                        p = doc.add_paragraph(f"[Gambar: {el.image_filename}]")
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        p.paragraph_format.first_line_indent = Inches(0)
                        has_any_content = True

                elif el.is_text():
                    if in_image_section:
                        sec = doc.add_section(WD_SECTION_START.NEW_PAGE)
                        _apply_text_section_format(sec)
                        in_image_section = False

                    has_any_content = True
                    translated_text = chunk_trans.get(i, el.text)

                    if el.element_type == ElementType.HEADING:
                        level = min(max(el.heading_level or 1, 1), 4)
                        h = doc.add_heading(translated_text, level=level)
                        h.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        h.paragraph_format.space_before = Pt(20)
                        h.paragraph_format.space_after = Pt(10)
                        h.paragraph_format.first_line_indent = Inches(0)
                        for r in h.runs:
                            r.font.bold = True
                            if level == 1:
                                r.font.size = Pt(18)
                            elif level == 2:
                                r.font.size = Pt(16)
                            elif level == 3:
                                r.font.size = Pt(14)

                    else:
                        if bilingual and el.text.strip() and el.text.strip() != translated_text.strip():
                            p_orig = doc.add_paragraph(el.text)
                            p_orig.paragraph_format.space_after = Pt(2)
                            for r in p_orig.runs:
                                r.italic = True

                        p = doc.add_paragraph(translated_text)

                        # Deteksi pembatas scene/bab (seperti ***, ---, ◇◇◇)
                        has_alnum = any(c.isalnum() for c in translated_text)
                        if not has_alnum and 0 < len(translated_text.strip()) < 20:
                            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            p.paragraph_format.space_before = Pt(12)
                            p.paragraph_format.space_after = Pt(12)
                            p.paragraph_format.first_line_indent = Inches(0)

        # Pastikan direktori output ada
        out_dir = os.path.dirname(os.path.abspath(output_file_path))
        if out_dir:
            os.makedirs(out_dir, exist_ok=True)

        doc.save(output_file_path)
