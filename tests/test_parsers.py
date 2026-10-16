import unittest
import os
import tempfile
import docx
from src.parsers import get_parser_for_file, ElementType
from src.builder import get_builder_for_file
from src.translator.checkpoint import CheckpointManager

class TestGeminiTranslatorCore(unittest.TestCase):
    def setUp(self):
        self.test_dir = tempfile.mkdtemp()
        self.sample_docx = os.path.join(self.test_dir, "sample.docx")

        # Buat file DOCX sampel untuk pengujian
        doc = docx.Document()
        doc.add_heading("Bab 1: Petualangan Baru", level=1)
        doc.add_paragraph("Ini adalah baris pertama dari novel sampel yang digunakan untuk pengujian unit parser.")
        doc.add_paragraph("Baris kedua berisi teks dialog: 'Kita harus pergi sekarang!' seru sang pahlawan.")
        doc.save(self.sample_docx)

    def test_docx_parsing(self):
        parser = get_parser_for_file(self.sample_docx, max_words_per_chunk=100)
        chunks = parser.parse(self.sample_docx)

        self.assertGreater(len(chunks), 0)
        first_chunk = chunks[0]
        self.assertGreater(len(first_chunk.elements), 0)

        # Pastikan ada heading dan paragraf
        types = [el.element_type for el in first_chunk.elements]
        self.assertIn(ElementType.HEADING, types)
        self.assertIn(ElementType.PARAGRAPH, types)

    def test_checkpoint_manager(self):
        checkpoint = CheckpointManager(self.sample_docx, cache_dir=os.path.join(self.test_dir, ".cache"))
        chunk_id = 0
        translated_data = {0: "Bab 1: Petualangan Baru (Terjemahan)"}

        # Simpan chunk
        checkpoint.save_chunk(chunk_id, translated_data)
        self.assertTrue(checkpoint.is_chunk_translated(chunk_id))

        # Ambil kembali
        loaded = checkpoint.get_chunk_translation(chunk_id)
        self.assertEqual(loaded[0], "Bab 1: Petualangan Baru (Terjemahan)")

    def test_docx_builder(self):
        output_docx = os.path.join(self.test_dir, "output.docx")
        parser = get_parser_for_file(self.sample_docx)
        chunks = parser.parse(self.sample_docx)

        dummy_translations = {
            0: {
                0: "Bab 1: Petualangan Baru (ID)",
                1: "Ini adalah terjemahan baris pertama.",
                2: "Ini adalah terjemahan baris kedua."
            }
        }

        builder = get_builder_for_file(output_docx)
        builder.build(chunks, dummy_translations, output_docx)

        self.assertTrue(os.path.exists(output_docx))

if __name__ == "__main__":
    unittest.main()
